import io
import math
import re
import zipfile
from copy import copy
from xml.etree import ElementTree as ET
 
import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties
 
TARGET_SHEET_KEYWORDS = [
    ("新規取得_年初", ["①", "新規取得", "年初"]),
    ("新規取得_追加", ["①", "新規取得", "追加用"]),
    ("補修リース_年初", ["②", "補修", "年初"]),
    ("補修リース_追加", ["②", "補修", "追加用"]),
    ("前年度追加", ["⓷", "新規取得", "追加"]),
]
 
# 明細は9行目から。1ページ分の行数・合計行の位置はテンプレートから自動判定する
BODY_START_ROW = 9
KEEP_SHEETS = ["償却率", "種類", "➃該当なし"]
 
# --- XML名前空間（プルダウン読み取り用）---
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
NS_XM = "http://schemas.microsoft.com/office/excel/2006/main"
 
 
# =====================================================================
# 共通ヘルパー
# =====================================================================
def find_matching_sheet(sheet_list, keywords):
    for s in sheet_list:
        if all(kw in s for kw in keywords):
            return s
    return None
 
 
def find_dept_col(df):
    """部署（施設）列を推定する。"""
    for col in df.columns:
        if "部署" in str(col) or "工場" in str(col) or "冷蔵" in str(col):
            return col
    if len(df.columns) >= 2:
        return df.columns[1]
    return None
 
 
def _truthy(v):
    return str(v).lower() in ("1", "true")
 
 
# =====================================================================
# プルダウン（入力規則）の読み取り
#   openpyxl は「他シート参照のプルダウン（x14拡張）」を読み込み時に捨ててしまうため、
#   元ファイルのXMLから直接読み取って、通常形式の入力規則として書き戻す。
# =====================================================================
def read_validation_rules(base_bytes):
    """{シート名: [ルール, ...]} を返す。通常形式とx14拡張形式の両方を拾う。"""
    z = zipfile.ZipFile(io.BytesIO(base_bytes))
    wb_root = ET.fromstring(z.read("xl/workbook.xml"))
    rel_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rel_root}
 
    rules_by_sheet = {}
    for sh in wb_root.iter(f"{{{NS_MAIN}}}sheet"):
        target = targets.get(sh.get(f"{{{NS_REL}}}id"))
        if not target:
            continue
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        try:
            root = ET.fromstring(z.read(path))
        except KeyError:
            continue
 
        rules = []
        # 通常形式
        for dv in root.iter(f"{{{NS_MAIN}}}dataValidation"):
            rules.append(
                _make_rule(
                    dv,
                    dv.findtext(f"{{{NS_MAIN}}}formula1"),
                    dv.findtext(f"{{{NS_MAIN}}}formula2"),
                    dv.get("sqref"),
                )
            )
        # x14拡張形式（別シート参照のプルダウンはこちらに入っている）
        for dv in root.iter(f"{{{NS_X14}}}dataValidation"):
            rules.append(
                _make_rule(
                    dv,
                    dv.findtext(f"{{{NS_X14}}}formula1/{{{NS_XM}}}f"),
                    dv.findtext(f"{{{NS_X14}}}formula2/{{{NS_XM}}}f"),
                    dv.findtext(f"{{{NS_XM}}}sqref"),
                )
            )
        rules_by_sheet[sh.get("name")] = [r for r in rules if r]
    return rules_by_sheet
 
 
def _make_rule(dv, formula1, formula2, sqref):
    dv_type = dv.get("type")
    if not dv_type or dv_type == "none" or not sqref:
        return None
    return {
        "type": dv_type,
        "operator": dv.get("operator"),
        "formula1": formula1,
        "formula2": formula2,
        "allow_blank": _truthy(dv.get("allowBlank")),
        "show_error": _truthy(dv.get("showErrorMessage")),
        "show_input": _truthy(dv.get("showInputMessage")),
        "error_title": dv.get("errorTitle"),
        "error": dv.get("error"),
        "prompt_title": dv.get("promptTitle"),
        "prompt": dv.get("prompt"),
        "ranges": sqref.split(),
    }
 
 
def _remap_range(rng, start_row, old_end_row, new_end_row):
    """明細範囲（9〜28行）全体にかかっている範囲は、全ページ分（9〜最終行）に広げる。"""
    min_c, min_r, max_c, max_r = range_boundaries(rng)
    if min_r == start_row and max_r == old_end_row:
        max_r = new_end_row
    a = f"{get_column_letter(min_c)}{min_r}"
    b = f"{get_column_letter(max_c)}{max_r}"
    return a if a == b else f"{a}:{b}"
 
 
def _covers(ranges, cell):
    col = openpyxl.utils.cell.coordinate_from_string(cell)
    c = openpyxl.utils.column_index_from_string(col[0])
    r = col[1]
    for rng in ranges:
        min_c, min_r, max_c, max_r = range_boundaries(rng)
        if min_c <= c <= max_c and min_r <= r <= max_r:
            return True
    return False
 
 
def apply_validations(ws, rules, start_row, old_end_row, new_end_row):
    ws.data_validations.dataValidation = []
    all_ranges = []
    for r in rules:
        ranges = [
            _remap_range(x, start_row, old_end_row, new_end_row)
            for x in r["ranges"]
        ]
        all_ranges.extend(ranges)
        dv = DataValidation(
            type=r["type"],
            operator=r["operator"],
            formula1=r["formula1"],
            formula2=r["formula2"],
            allow_blank=r["allow_blank"],
            showErrorMessage=r["show_error"],
            showInputMessage=r["show_input"],
            errorTitle=r["error_title"],
            error=r["error"],
            promptTitle=r["prompt_title"],
            prompt=r["prompt"],
        )
        for rng in ranges:
            dv.add(rng)
        ws.add_data_validation(dv)
 
    # 単位セル(D2)にプルダウンが無いテンプレートの場合だけ、従来の固定リストを付ける
    if not _covers(all_ranges, "D2"):
        dv_unit = DataValidation(
            type="list", formula1='"千円,ドル,元,ルピー"', allow_blank=True
        )
        ws.add_data_validation(dv_unit)
        dv_unit.add("D2")
 
 
# =====================================================================
# シートへの書き込み（ページ自動追加）
# =====================================================================
def _find_total_row(ws, start_row):
    for r in range(start_row, ws.max_row + 1):
        if str(ws.cell(row=r, column=2).value).strip() == "合計":
            return r
    return start_row + 20  # 見つからなければ「20行/ページ」とみなす
 
 
def _print_area_cols(ws):
    """テンプレートの印刷範囲の列幅（開始列・終了列）を取得する。"""
    try:
        area = str(ws.print_area).split("!")[-1].split(",")[0].replace("$", "")
        min_c, _, max_c, _ = range_boundaries(area)
        return min_c, max_c
    except Exception:
        return 1, ws.max_column
 
 
def fill_sheet(ws, df, rules):
    """
    ws に df の内容を書き込む。
    ・1ページ = テンプレートの明細行数（通常20行）。件数が増えたらページ単位で増やす
    ・各ページに書式・行の高さ・No.・計算式をコピーし、改ページを入れる
    ・合計行は最終ページの末尾に置く
    ・プルダウンは全ページ分の範囲に適用する
    """
    start = BODY_START_ROW
    total_tmpl = _find_total_row(ws, start)
    rows_per_page = total_tmpl - start
    max_col = ws.max_column
    min_pc, max_pc = _print_area_cols(ws)
 
    # --- テンプレート行のスナップショット ---
    body_style = [
        [copy(ws.cell(row=r, column=c)._style) for c in range(1, max_col + 1)]
        for r in range(start, total_tmpl)
    ]
    body_height = [ws.row_dimensions[r].height for r in range(start, total_tmpl)]
    total_style = [
        copy(ws.cell(row=total_tmpl, column=c)._style) for c in range(1, max_col + 1)
    ]
    total_height = ws.row_dimensions[total_tmpl].height
    total_formulas = {
        c: ws.cell(row=total_tmpl, column=c).value
        for c in range(1, max_col + 1)
        if isinstance(ws.cell(row=total_tmpl, column=c).value, str)
        and ws.cell(row=total_tmpl, column=c).value.startswith("=")
    }
    # 明細1行目の計算式（償却方法・償却率・償却額・取得価額など）を各行に展開する
    row_formulas = {}
    for c in range(1, max_col + 1):
        v = ws.cell(row=start, column=c).value
        if isinstance(v, str) and v.startswith("=") and "#REF!" not in v:
            row_formulas[c] = v
    # No.の書式（「1,2,3…」か「追1,追2…」か）
    b_first = ws.cell(row=start, column=2).value
    prefix = re.sub(r"\d+$", "", b_first) if isinstance(b_first, str) else ""
 
    # --- ページ数の決定 ---
    n = len(df)
    pages = max(1, math.ceil(n / rows_per_page))
    body_rows = pages * rows_per_page
    last_body = start + body_rows - 1
    total_row = last_body + 1
 
    # --- 既存の値をクリア ---
    for r in range(start, max(ws.max_row, total_row) + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).value = None
 
    # 単位
    if not ws.cell(row=2, column=4).value:
        ws.cell(row=2, column=4, value="千円")
 
    # --- 明細（全ページ分。余った行は空欄のまま書式・No.・計算式を維持）---
    records = df.reset_index(drop=True)
    for i in range(body_rows):
        r = start + i
        tmpl_idx = i % rows_per_page
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c)._style = copy(body_style[tmpl_idx][c - 1])
        if body_height[tmpl_idx] is not None:
            ws.row_dimensions[r].height = body_height[tmpl_idx]
 
        ws.cell(row=r, column=2, value=(f"{prefix}{i + 1}" if prefix else i + 1))
        for c, formula in row_formulas.items():
            ws.cell(row=r, column=c).value = Translator(
                formula, origin=f"{get_column_letter(c)}{start}"
            ).translate_formula(f"{get_column_letter(c)}{r}")
 
        if i < n:
            for c_idx, val in enumerate(records.iloc[i]):
                col = c_idx + 1
                if col == 2 or col in row_formulas or col > max_col:
                    continue
                if pd.notna(val):
                    ws.cell(row=r, column=col, value=val)
 
    # --- 合計行（最終ページの末尾）---
    for c in range(1, max_col + 1):
        ws.cell(row=total_row, column=c)._style = copy(total_style[c - 1])
    if total_height is not None:
        ws.row_dimensions[total_row].height = total_height
    ws.cell(row=total_row, column=2, value="合計")
    pat = re.compile(r"([A-Z]{1,3})%d:([A-Z]{1,3})%d" % (start, total_tmpl - 1))
    for c, formula in total_formulas.items():
        ws.cell(row=total_row, column=c).value = pat.sub(
            lambda m: f"{m.group(1)}{start}:{m.group(2)}{last_body}", formula
        )
 
    # --- プルダウン（全ページ分に適用）---
    apply_validations(ws, rules, start, total_tmpl - 1, last_body)
 
    # --- 印刷設定：ページ数に応じて印刷範囲と改ページを設定 ---
    ws.print_area = (
        f"{get_column_letter(min_pc)}1:{get_column_letter(max_pc)}{total_row}"
    )
    ws.print_title_rows = "1:8"  # 2ページ目以降も見出し(1〜8行)を繰り返す
    ws.row_breaks = RowBreak()
    for p in range(1, pages):
        ws.row_breaks.append(Break(id=start + p * rows_per_page - 1))
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.scale = None
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1  # 横は1ページに収める
    ws.page_setup.fitToHeight = 0  # 縦は自動（改ページプレビューで ページ1,2,3… となる）
 
    return n, pages
 
 
def _sheet_is_empty_legacy(ws):
    """集約対象外シート用の従来判定：9行目以降のB列に有効な値が無ければ空。"""
    return not any(
        str(ws.cell(row=r, column=2).value).strip()
        not in ["None", "nan", "合計", "計", "No."]
        for r in range(9, ws.max_row + 1)
    )
 
 
def build_workbook(base_bytes, combined_rows, facilities=None):
    """
    集約データからExcelを作る。facilities が None なら全国版、
    リスト（空でも可）なら該当施設だけを抽出した版。
    """
    rules_by_sheet = read_validation_rules(base_bytes)
    wb = openpyxl.load_workbook(io.BytesIO(base_bytes))
    has_data = {}
 
    for key, keywords in TARGET_SHEET_KEYWORDS:
        matched = find_matching_sheet(wb.sheetnames, keywords)
        if not matched:
            continue
        df_list = combined_rows.get(key, [])
        if df_list:
            df = pd.concat(df_list, ignore_index=True)
        else:
            df = pd.DataFrame()
 
        if facilities is not None:
            dept_col = find_dept_col(df) if not df.empty else None
            if dept_col is not None and facilities:
                df = df[
                    df[dept_col].astype(str).str.strip().isin(facilities)
                ].copy()
            else:
                df = df.iloc[0:0]
 
        fill_sheet(wb[matched], df, rules_by_sheet.get(matched, []))
        has_data[matched] = len(df) > 0
 
    # --- 不要シート削除（データが無い追加用シートなど）---
    to_remove = []
    for s in wb.sheetnames:
        if s in KEEP_SHEETS:
            continue
        if not ("追加" in s or "年間減価償却費" in s or "償却月数" in s):
            continue
        empty = (not has_data[s]) if s in has_data else _sheet_is_empty_legacy(wb[s])
        if empty:
            to_remove.append(s)
    for s in to_remove:
        wb.remove(wb[s])
 
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
 
 
# =====================================================================
# 集約
# =====================================================================
def aggregate_files(uploaded_files):
    combined_rows = {key: [] for key, _ in TARGET_SHEET_KEYWORDS}
    processed_count = 0
    errors = []
 
    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        if (
            "★" in filename
            or "全国" in filename
            or "集約" in filename
            or "0000年度" in filename
        ):
            continue
        try:
            xl = pd.ExcelFile(uploaded_file)
            for key, keywords in TARGET_SHEET_KEYWORDS:
                matched_sheet = find_matching_sheet(xl.sheet_names, keywords)
                if not matched_sheet:
                    continue
                df = xl.parse(matched_sheet, header=7)
                if "資産名称" in df.columns:
                    df_valid = df.dropna(subset=["資産名称"]).copy()
                    if "No." in df_valid.columns:
                        df_valid = df_valid[df_valid["No."].astype(str) != "合計"]
                    if not df_valid.empty:
                        combined_rows[key].append(df_valid)
            processed_count += 1
        except Exception as e:
            errors.append(f"❌ {filename} の読み込みエラー: {e}")
 
    return combined_rows, processed_count, errors
 
 
# =====================================================================
# 画面
# =====================================================================
def main():
    st.set_page_config(
        page_title="設備投資計画フォーム 自動集約＆施設抽出ツール",
        page_icon="📊",
        layout="wide",
    )
 
    st.title("📊 設備投資計画フォーム 自動集約＆施設抽出ツール")
    st.markdown("""
各部署から回収した「設備投資計画フォーム（Excel）」をまとめて集約します。  
画面上のプルダウンから施設（部署）を複数選択して、全国版と冷蔵部版のファイルをいつでも両方ダウンロードできます。
""")
 
    st.sidebar.header("📂 設定メニュー")
    template_file = st.sidebar.file_uploader(
        "【任意】まとめ用テンプレート（0000年度/全国まとめフォーム）",
        type=["xlsx"],
    )
 
    uploaded_files = st.file_uploader(
        "各部署のExcelファイルをここにドラッグ＆ドロップ（複数可）",
        type=["xlsx"],
        accept_multiple_files=True,
    )
 
    if uploaded_files:
        st.info(f"📁 現在 {len(uploaded_files)} 個のファイルが選択されています。")
 
        if st.button("⚡ データを集約する", type="primary"):
            combined_rows, processed_count, errors = aggregate_files(uploaded_files)
            for msg in errors:
                st.error(msg)
 
            st.session_state["combined_rows"] = combined_rows
            st.session_state["has_aggregated"] = True
            # 新規集約時は古いダウンロード用データをクリア
            st.session_state.pop("output_nat_bytes", None)
            st.session_state.pop("output_dep_bytes", None)
            st.success(
                f"🎉 {processed_count} 件の部署データ集約が完了しました！"
                "下に施設の選択が表示されます。"
            )
 
    # --- 集計完了後のUI表示 ---
    if st.session_state.get("has_aggregated", False):
        combined_rows = st.session_state["combined_rows"]
 
        # 全データの中から施設名・部署名のユニークなリストを取得
        all_facilities = set()
        for df_list in combined_rows.values():
            for df in df_list:
                dept_col = find_dept_col(df)
                if dept_col is not None and dept_col in df.columns:
                    facs = (
                        df[dept_col].dropna().astype(str).str.strip().unique().tolist()
                    )
                    all_facilities.update([f for f in facs if f != "nan"])
 
        sorted_facilities = sorted(all_facilities)
 
        st.divider()
        st.subheader("🏢 別ファイル（冷蔵部等）作成用の施設選択")
 
        default_selected = [f for f in sorted_facilities if "冷蔵" in f]
 
        dept_file_label = st.text_input(
            "抽出ファイル名（部署名）",
            value="冷蔵部",
            help="例：「冷蔵部」と入力すると『★【冷蔵部】設備投資計画集約版.xlsx』が出力されます。",
        )
 
        selected_facilities = st.multiselect(
            "抽出対象の施設・部署を選択してください（複数選択可）：",
            options=sorted_facilities,
            default=default_selected,
        )
 
        if st.button("📄 両方のExcelファイルを一括生成する", type="primary"):
            base_file = template_file if template_file else uploaded_files[0]
            try:
                base_bytes = base_file.getvalue()
 
                st.session_state["output_nat_bytes"] = build_workbook(
                    base_bytes, combined_rows, facilities=None
                )
                st.session_state["output_dep_bytes"] = build_workbook(
                    base_bytes, combined_rows, facilities=selected_facilities
                )
                st.session_state["dept_file_label"] = dept_file_label
 
                st.success(
                    "🎉 両方のファイルが正常に生成されました！いつでもダウンロード可能です。"
                )
            except Exception as e:
                st.error(f"ファイル生成中にエラーが発生しました: {e}")
 
        # セッションにデータが存在していれば、ダウンロードボタンを常時表示
        if (
            "output_nat_bytes" in st.session_state
            and "output_dep_bytes" in st.session_state
        ):
            st.divider()
            st.markdown("### 📥 ダウンロードセンター")
            label_name = st.session_state.get("dept_file_label", "冷蔵部")
            xlsx_mime = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
 
            col1, col2 = st.columns(2)
            with col1:
                st.download_button(
                    label="📥 【全国版】集約ファイルをダウンロード",
                    data=st.session_state["output_nat_bytes"],
                    file_name="★【全国】設備投資計画集約版.xlsx",
                    mime=xlsx_mime,
                    key="download_nat",
                )
            with col2:
                st.download_button(
                    label=f"📥 【{label_name}版】抽出ファイルをダウンロード",
                    data=st.session_state["output_dep_bytes"],
                    file_name=f"★【{label_name}】設備投資計画集約版.xlsx",
                    mime=xlsx_mime,
                    key="download_dep",
                )
 
 
if __name__ == "__main__":
    main()
 
